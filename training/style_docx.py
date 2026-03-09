"""Post-process pandoc Word docs for professional styling."""

import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ACCENT = RGBColor(0x00, 0x7A, 0xCC)
DARK = RGBColor(0x2D, 0x2D, 0x2D)
MUTED = RGBColor(0x66, 0x66, 0x66)


def _set_runs(paragraph, size, color, bold=False, italic=False, font_name="Calibri"):
    """Apply font properties to all runs in a paragraph."""
    for run in paragraph.runs:
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = font_name
        if bold:
            run.bold = True
        if italic:
            run.font.italic = True


def _set_cell_shading(cell, fill_color):
    """Apply background shading to a table cell."""
    tc_pr = cell._element.get_or_add_tcPr()
    shading_elem = tc_pr.makeelement(
        qn("w:shd"),
        {
            qn("w:val"): "clear",
            qn("w:color"): "auto",
            qn("w:fill"): fill_color,
        },
    )
    tc_pr.append(shading_elem)


def _style_paragraphs(doc):
    """Style headings and body paragraphs."""
    for paragraph in doc.paragraphs:
        style_name = paragraph.style.name

        if style_name == "Title":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_runs(paragraph, 28, ACCENT, bold=True)

        elif style_name == "Subtitle":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _set_runs(paragraph, 16, MUTED)

        elif "Heading 1" in style_name:
            _set_runs(paragraph, 22, ACCENT, bold=True)
            paragraph.paragraph_format.space_before = Pt(24)
            paragraph.paragraph_format.space_after = Pt(8)

        elif "Heading 2" in style_name:
            _set_runs(paragraph, 16, DARK, bold=True)
            paragraph.paragraph_format.space_before = Pt(18)
            paragraph.paragraph_format.space_after = Pt(6)

        elif "Heading 3" in style_name:
            _set_runs(paragraph, 13, ACCENT, bold=True)
            paragraph.paragraph_format.space_before = Pt(14)
            paragraph.paragraph_format.space_after = Pt(4)

        elif style_name.startswith("Body") or style_name == "Normal":
            for run in paragraph.runs:
                if run.font.size is None:
                    run.font.size = Pt(11)
                if run.font.name is None or run.font.name == "":
                    run.font.name = "Calibri"
                if run.font.color.rgb is None:
                    run.font.color.rgb = DARK

        # Block quotes (used for > notes)
        elif "Block" in style_name or "Quote" in style_name:
            _set_runs(paragraph, 10, MUTED, italic=True)


def _style_tables(doc):
    """Style tables with header row and alternating row shading."""
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        if not table.rows:
            continue

        # Header row
        for cell in table.rows[0].cells:
            _set_cell_shading(cell, "007ACC")
            for paragraph in cell.paragraphs:
                _set_runs(paragraph, 10, RGBColor(0xFF, 0xFF, 0xFF), bold=True)

        # Data rows
        for i, row in enumerate(table.rows[1:], 1):
            if i % 2 == 0:
                for cell in row.cells:
                    _set_cell_shading(cell, "F0F6FC")

            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
                        run.font.name = "Calibri"


def style_document(filepath):
    doc = Document(filepath)

    # -- Page margins --
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    _style_paragraphs(doc)
    _style_tables(doc)

    doc.save(filepath)
    print(f"Styled: {filepath}")


if __name__ == "__main__":
    files = (
        sys.argv[1:]
        if len(sys.argv) > 1
        else [
            "LAB-GUIDE.docx",
            "REFERENCE.docx",
            "SETUP.docx",
        ]
    )
    for f in files:
        style_document(f)
